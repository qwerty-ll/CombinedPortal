import io
from datetime import date, datetime

import pytest
from docx import Document

from app.services import documents, document_drafts, eios, timetable
from conftest import CSRF, login_student

NOW = datetime(2026, 9, 24, 10, 15, tzinfo=timetable.MSK)  # Thursday


def lesson(day, start, end, discipline, room="Б-305", teacher="Иванов И.И."):
    return {"дата": f"{day}T00:00:00", "начало": start, "конец": end, "дисциплина": discipline, "аудитория": room,
            "группа": "24-ИСбо-1", "преподаватель": teacher, "номерПодгруппы": 0, "замена": False}


ANSWERS = {
    ("raspGrouplist", frozenset({"year": "2026-2027"}.items())): {"state": 1, "data": [{"id": 4242, "name": "24-ИСбо-1"}]},
    ("Rasp", frozenset({"year": "2026-2027", "idGroup": 4242}.items())): {"state": 1, "data": {"rasp": [
        lesson("2026-09-23", "08:30", "10:00", "лек Философия"),
        lesson("2026-09-23", "10:10", "11:40", "пр Базы данных", room="Б-407", teacher="Сидоров С.С."),
        lesson("2026-09-24", "08:30", "10:00", "пр Философия", teacher="Петрова П.П."),
    ]}},
}


@pytest.fixture
def student(app, fake_eios, monkeypatch):
    async def fetch_json(endpoint, params, timeout=5.0):
        return ANSWERS.get((endpoint, frozenset(params.items())))

    monkeypatch.setattr(eios, "fetch_json", fetch_json)
    monkeypatch.setattr(timetable, "msk_now", lambda: NOW)
    return login_student(app, fake_eios, full_name="Иванова Анна Сергеевна", group="24-ИСбо-1")


def docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)


EXPLANATORY = {
    "full_name": "Иванова Анна Сергеевна", "group": "24-ИСбо-1", "course": 3,
    "date_from": "2026-09-23", "reason": "Болезнь", "attachment": "справка из поликлиники",
    "pairs": [
        {"start": "08:30", "end": "10:00", "discipline": "Философия", "kind": "лекция", "teacher": "Иванов И.И."},
        {"start": "10:10", "end": "11:40", "discipline": "Базы данных", "kind": "практика", "teacher": "Сидоров С.С."},
    ],
}


def test_person_is_declined_for_the_address():
    anna = documents.Person("Иванова Анна Сергеевна", "24-ИСбо-1")
    assert anna.genitive == "Ивановой Анны Сергеевны" and anna.signature == "А. С. Иванова"
    assert anna.word("студент", "студентка", "?") == "студентка"
    assert documents.Person("Барило Илья Иванович", "x").genitive == "Барило Ильи Ивановича"
    assert documents.Person("Шевченко Мария Петровна", "x").genitive == "Шевченко Марии Петровны"
    assert documents.course_of("24-ИСбо-1", date(2026, 9, 24)) == 3
    assert documents.course_of("24-ИСбо-1", date(2025, 3, 1)) == 1
    assert documents.course_of("Деканат ИВИТШ", date(2026, 9, 24)) is None


def test_date_ranges():
    assert documents.range_text(date(2026, 9, 24), None) == "24 сентября 2026 г."
    assert documents.range_text(date(2026, 9, 21), date(2026, 9, 24)) == "с 21 по 24 сентября 2026 г."
    assert documents.range_text(date(2026, 9, 30), date(2026, 10, 2)) == "с 30 сентября по 2 октября 2026 г."
    assert documents.range_text(date(2025, 12, 30), date(2026, 1, 12)) == "с 30 декабря 2025 г. по 12 января 2026 г."


@pytest.mark.parametrize("q, expected", [
    ("объяснительная за вчера", date(2026, 9, 23)),
    ("позавчера не был", date(2026, 9, 22)),
    ("в понедельник болел", date(2026, 9, 21)),
    ("в четверг", date(2026, 9, 24)),
    ("в пятницу", date(2026, 9, 18)),
    ("пропустил 15.09", date(2026, 9, 15)),
    ("пропустил 3 сентября", date(2026, 9, 3)),
    ("24.12", date(2025, 12, 24)),
    ("просто объяснительная", None),
])
def test_parse_past_day(q, expected):
    assert document_drafts.parse_past_day(q, date(2026, 9, 24)) == expected


def test_documents_need_a_signed_in_student(client):
    assert client.post("/api/v1/documents/explanatory", json=EXPLANATORY, headers=CSRF).status_code == 401
    assert client.get("/api/v1/documents/pairs", params={"date": "2026-09-23"}).status_code == 401


def test_explanatory_note_as_word(student):
    r = student.post("/api/v1/documents/explanatory", params={"format": "docx"}, json=EXPLANATORY, headers=CSRF)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert "filename*=UTF-8''%D0%9E%D0%B1%D1%8A" in r.headers["content-disposition"]  # "Объяснительная_…"
    text = docx_text(r.content)
    assert "Директору Высшей ИТ-школы КГУ\n__________________________\nстудентки 3 курса группы 24-ИСбо-1\nИвановой Анны Сергеевны" in text
    assert ("Я, Иванова Анна Сергеевна, студентка группы 24-ИСбо-1, отсутствовала на занятиях "
            "23 сентября 2026 г. по причине: болезнь.") in text
    assert "– 08:30–10:00 — Философия (лекция), преподаватель Иванов И.И.;" in text
    assert "– 10:10–11:40 — Базы данных (практика), преподаватель Сидоров С.С." in text
    assert "Подтверждающий документ прилагаю: справка из поликлиники." in text
    assert text.endswith("«24» сентября 2026 г.\t____________ / А. С. Иванова")
    style = Document(io.BytesIO(r.content)).styles["Normal"].font
    assert style.name == "Times New Roman" and style.size.pt == 14


def test_explanatory_note_as_pdf(student):
    r = student.post("/api/v1/documents/explanatory", params={"format": "pdf"}, json=EXPLANATORY, headers=CSRF)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and b"LiberationSerif" in r.content


def test_a_period_of_absence_lists_no_pairs(student):
    body = {**EXPLANATORY, "date_to": "2026-09-25", "attachment": ""}
    text = docx_text(student.post("/api/v1/documents/explanatory", json=body, headers=CSRF).content)
    assert "на занятиях с 23 по 25 сентября 2026 г. по причине: болезнь." in text
    assert "Пропущенные занятия" not in text and "Подтверждающий" not in text


def test_retake_request(student):
    body = {"full_name": "Тестов Тест Тестович", "group": "24-ИСбо-1", "course": 3, "discipline": "Базы данных",
            "control": "экзамен", "teacher": "Сидоров С.С.", "reason": documents.RETAKE_REASONS[0]}
    text = docx_text(student.post("/api/v1/documents/retake", json=body, headers=CSRF).content)
    assert "студента 3 курса группы 24-ИСбо-1\nТестова Теста Тестовича\nЗаявление" in text
    assert ("Прошу разрешить мне пересдачу экзамена по дисциплине «Базы данных» (преподаватель — Сидоров С.С.) "
            "в связи с получением неудовлетворительной оценки.") in text


@pytest.mark.parametrize("change", [
    {"date_to": "2026-09-20"},
    {"full_name": "Иванова"},
    {"full_name": "<script>alert(1)</script> x"},
    {"reason": "x" * 400},
    {"pairs": [{"start": "8:30:00", "end": "10:00", "discipline": "X"}]},
])
def test_explanatory_validation(student, change):
    assert student.post("/api/v1/documents/explanatory", json={**EXPLANATORY, **change}, headers=CSRF).status_code == 422


def test_pairs_of_a_day(student):
    pairs = student.get("/api/v1/documents/pairs", params={"date": "2026-09-23"}).json()["pairs"]
    assert [(p["start"], p["discipline"], p["teacher"]) for p in pairs] == [
        ("08:30", "Философия", "Иванов И.И."), ("10:10", "Базы данных", "Сидоров С.С."),
    ]


def ask(client, message):
    r = client.post("/api/v1/chat", json={"message": message}, headers=CSRF)
    assert r.status_code == 200, r.text
    return r.json()


def test_the_cat_drafts_an_explanatory_note(student):
    data = ask(student, "Напиши объяснительную, вчера болел")
    assert data["reply"] == "Собрал объяснительную за вчера: 2 пары из расписания. Причина — болезнь. Проверь и скачай."
    draft = data["document"]
    assert draft["kind"] == "explanatory"
    fields = draft["fields"]
    assert (fields["full_name"], fields["group"], fields["course"]) == ("Иванова Анна Сергеевна", "24-ИСбо-1", 3)
    assert fields["date_from"] == "2026-09-23" and fields["reason"] == "болезнь"
    assert [p["discipline"] for p in fields["pairs"]] == ["Философия", "Базы данных"] and all(p["checked"] for p in fields["pairs"])
    # No reason given: the cat asks for it
    assert ask(student, "нужна объяснительная за 22.09")["reply"] == (
        "Собрал объяснительную за 22 сентября. Пар в расписании на этот день не нашёл — проверь дату. "
        "Укажи причину — и можно скачивать."
    )


def test_the_cat_drafts_a_retake_request(student):
    data = ask(student, "Хочу пересдать философию")
    assert data["reply"] == ("Собрал заявление на пересдачу экзамена по дисциплине «Философия» "
                             "(преподаватель — Иванов И.И.). Проверь и скачай.")
    fields = data["document"]["fields"]
    assert (fields["discipline"], fields["teacher"], fields["control"]) == ("Философия", "Иванов И.И.", "экзамен")
    assert fields["reason"] == documents.RETAKE_REASONS[0]
    assert {o["discipline"] for o in data["document"]["options"]["disciplines"]} == {"Философия", "Базы данных"}
    fields = ask(student, "заявление на пересдачу зачёта, я болел")["document"]["fields"]
    assert fields["discipline"] == "" and fields["control"] == "зачёт" and fields["reason"] == documents.RETAKE_REASONS[1]


def test_guests_are_asked_to_sign_in(client):
    data = ask(client, "напиши объяснительную")
    assert data["document"] is None and data["actions"] == [{"label": "Войти через ЭИОС", "to": "/profile"}]
