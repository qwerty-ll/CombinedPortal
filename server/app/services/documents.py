"""Paperwork ВИТШик fills in: an explanatory note for missed classes and a request to retake an exam.

One layout is rendered twice: as DOCX to edit and print (Times New Roman 14, as university offices expect)
and as PDF to print as is (Liberation Serif, metrically identical to Times New Roman, bundled with the app).
"""
import io
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import List, Optional
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pytrovich.detector import PetrovichGenderDetector
from pytrovich.enums import Case, Gender, NamePart
from pytrovich.maker import PetrovichDeclinationMaker
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import settings

MONTHS_GENITIVE = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]
CONTROL_GENITIVE = {
    "экзамен": "экзамена",
    "зачёт": "зачёта",
    "дифференцированный зачёт": "дифференцированного зачёта",
}
EXPLANATORY_REASONS = ["болезнь", "семейные обстоятельства", "участие в мероприятии университета"]
RETAKE_REASONS = [
    "получением неудовлетворительной оценки",
    "неявкой на промежуточную аттестацию по уважительной причине",
]

_FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"
pdfmetrics.registerFont(TTFont("DocSerif", str(_FONTS / "LiberationSerif-Regular.ttf")))
pdfmetrics.registerFont(TTFont("DocSerif-Bold", str(_FONTS / "LiberationSerif-Bold.ttf")))

_declension = PetrovichDeclinationMaker()
_gender_detector = PetrovichGenderDetector()

# Page geometry of Russian office paperwork: A4, margins 3 / 1.5 / 2 / 2 cm
_MARGIN_LEFT, _MARGIN_RIGHT, _MARGIN_Y = 3.0, 1.5, 2.0
_TEXT_WIDTH = 21.0 - _MARGIN_LEFT - _MARGIN_RIGHT
_LIST_INDENT, _DASH_HANG = 1.75, 0.5
_DASH = "–\u00a0"  # a list dash kept on the same line as its text


def _address_indent(lines: List[str]) -> float:
    """The "to / from" block sits on the right, as wide as its longest line (Liberation Serif has the
    metrics of Times New Roman, so the same indent fits the DOCX too)."""
    widest = max(pdfmetrics.stringWidth(line, "DocSerif", 14) for line in lines) / cm
    # 0.8 cm: the PDF frame keeps 6 pt of padding on each side, and a little air
    return min(9.0, max(5.5, _TEXT_WIDTH - widest - 0.8))


# --- The student -------------------------------------------------------------------------------

def clean_text(value: str, limit: int) -> str:
    """One line of plain text: no control characters, single spaces, capped length."""
    value = re.sub(r"[\x00-\x1f\x7f]", " ", value or "")
    return " ".join(value.split())[:limit]


def course_of(group: str, today: date) -> Optional[int]:
    """"24-ИСбо-1" in the 2026-2027 academic year → 3."""
    m = re.match(r"^(\d{2})-", group or "")
    if not m:
        return None
    start = today.year if today.month >= 9 else today.year - 1
    course = start - (2000 + int(m.group(1))) + 1
    return course if 1 <= course <= 6 else None


@dataclass
class Person:
    full_name: str
    group: str
    course: Optional[int] = None

    @property
    def parts(self) -> List[str]:
        return self.full_name.split()

    @property
    def gender(self) -> Optional[Gender]:
        parts = self.parts
        if len(parts) >= 3 and parts[2].lower().endswith(("вна", "чна", "кызы")):
            return Gender.FEMALE
        if len(parts) >= 3 and parts[2].lower().endswith(("вич", "ич", "оглы")):
            return Gender.MALE
        try:
            detected = _gender_detector.detect(firstname=parts[1] if len(parts) > 1 else None,
                                               middlename=parts[2] if len(parts) > 2 else None)
        except Exception:
            return None
        return detected if detected in (Gender.MALE, Gender.FEMALE) else None

    def word(self, male: str, female: str, unknown: str) -> str:
        return {Gender.MALE: male, Gender.FEMALE: female}.get(self.gender, unknown)

    @property
    def genitive(self) -> str:
        """"Иванов Иван Иванович" → "Иванова Ивана Ивановича"; left as is when it cannot be declined."""
        gender = self.gender
        if gender is None:
            return self.full_name
        kinds = [NamePart.LASTNAME, NamePart.FIRSTNAME, NamePart.MIDDLENAME]
        try:
            return " ".join(
                _declension.make(kind, gender, Case.GENITIVE, part) for kind, part in zip(kinds, self.parts)
            ) + ("" if len(self.parts) <= 3 else " " + " ".join(self.parts[3:]))
        except Exception:
            return self.full_name

    @property
    def signature(self) -> str:
        """"Иванов Иван Иванович" → "И. И. Иванов"."""
        parts = self.parts
        if len(parts) < 2:
            return self.full_name
        return " ".join(f"{p[0]}." for p in parts[1:3]) + f" {parts[0]}"


# --- Layout ------------------------------------------------------------------------------------

@dataclass
class Paper:
    """A one-page statement: address block, title, text, date and signature."""
    header: List[str]
    title: str
    paragraphs: List[str]
    items: List[str] = field(default_factory=list)  # a list after the first paragraph
    closing: List[str] = field(default_factory=list)
    dated: date = field(default_factory=date.today)
    signed: str = ""
    file_stem: str = "document"


def date_text(day: date) -> str:
    return f"{day.day} {MONTHS_GENITIVE[day.month - 1]} {day.year} г."


def range_text(first: date, last: Optional[date]) -> str:
    """"24 сентября 2026 г.", "с 21 по 24 сентября 2026 г.", "с 30 сентября по 2 октября 2026 г."."""
    if last is None or last == first:
        return date_text(first)
    if first.year == last.year and first.month == last.month:
        return f"с {first.day} по {last.day} {MONTHS_GENITIVE[last.month - 1]} {last.year} г."
    if first.year == last.year:
        return f"с {first.day} {MONTHS_GENITIVE[first.month - 1]} по {last.day} {MONTHS_GENITIVE[last.month - 1]} {last.year} г."
    return f"с {date_text(first)} по {date_text(last)}"


def _sentence_part(text: str) -> str:
    """User text in the middle of a sentence: no trailing period, lowercase start unless it is an abbreviation."""
    text = text.strip().rstrip(".;")
    if len(text) > 1 and text[0].isupper() and not text[1].isupper():
        text = text[0].lower() + text[1:]
    return text


def _address(person: Person) -> List[str]:
    student = person.word("студента", "студентки", "студента")
    course = f" {person.course} курса" if person.course else ""
    return [
        settings.DOCUMENT_ADDRESSEE_TITLE,
        settings.DOCUMENT_ADDRESSEE_NAME or "_" * 26,
        f"{student}{course} группы {person.group}",
        person.genitive,
    ]


@dataclass
class MissedPair:
    start: str
    end: str
    discipline: str
    kind: str = ""
    teacher: str = ""


def explanatory_note(person: Person, first: date, last: Optional[date], reason: str,
                     pairs: List[MissedPair], attachment: str, today: date) -> Paper:
    student = person.word("студент", "студентка", "студент(ка)")
    was_absent = person.word("отсутствовал", "отсутствовала", "отсутствовал(а)")
    paragraphs = [
        f"Я, {person.full_name}, {student} группы {person.group}, {was_absent} на занятиях "
        f"{range_text(first, last)} по причине: {_sentence_part(reason)}."
    ]
    items = []
    for i, pair in enumerate(pairs):
        kind = f" ({pair.kind})" if pair.kind else ""
        teacher = f", преподаватель {pair.teacher}" if pair.teacher else ""
        text = f"{pair.start}–{pair.end} — {pair.discipline}{kind}{teacher}".rstrip(";")
        last = i == len(pairs) - 1
        # "Иванов И.И." keeps the period of its initials: "…И.И.;" and, at the end, "…И.И."
        items.append(text if last and text.endswith(".") else text + ("." if last else ";"))
    if items:
        paragraphs.append("Пропущенные занятия:")
    closing = [f"Подтверждающий документ прилагаю: {_sentence_part(attachment)}."] if attachment.strip() else []
    return Paper(
        header=_address(person), title="Объяснительная записка", paragraphs=paragraphs, items=items,
        closing=closing, dated=today, signed=person.signature,
        file_stem=f"Объяснительная_{person.parts[0]}_{first:%d.%m.%Y}",
    )


def retake_request(person: Person, discipline: str, control: str, teacher: str, reason: str, today: date) -> Paper:
    teacher_part = f" (преподаватель — {teacher})" if teacher.strip() else ""
    text = (
        f"Прошу разрешить мне пересдачу {CONTROL_GENITIVE[control]} по дисциплине «{discipline.strip()}»"
        f"{teacher_part} в связи с {_sentence_part(reason)}."
    )
    return Paper(
        header=_address(person), title="Заявление", paragraphs=[text], dated=today, signed=person.signature,
        file_stem=f"Заявление_на_пересдачу_{person.parts[0]}_{today:%d.%m.%Y}",
    )


def _dated(day: date) -> str:
    return f"«{day.day:02d}» {MONTHS_GENITIVE[day.month - 1]} {day.year} г."


# --- DOCX --------------------------------------------------------------------------------------

def render_docx(paper: Paper) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(_MARGIN_LEFT), Cm(_MARGIN_RIGHT)
    section.top_margin = section.bottom_margin = Cm(_MARGIN_Y)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    fonts = normal.element.rPr.rFonts
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Times New Roman")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    indent = _address_indent(paper.header)
    for line in paper.header:
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Cm(indent)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(18)
    title.add_run(paper.title).bold = True

    def body(text: str, indent: bool = True):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.25)
        return p

    body(paper.paragraphs[0])
    for text in paper.paragraphs[1:]:
        body(text)
    for item in paper.items:
        # Left-aligned with a hanging dash: justified short lines would stretch
        p = body(_DASH + item, indent=False)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(_LIST_INDENT)
        p.paragraph_format.first_line_indent = Cm(-_DASH_HANG)
    for text in paper.closing:
        body(text)

    signature = doc.add_paragraph()
    signature.paragraph_format.space_before = Pt(42)
    signature.paragraph_format.tab_stops.add_tab_stop(Cm(_TEXT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    signature.add_run(f"{_dated(paper.dated)}\t____________ / {paper.signed}")

    doc.core_properties.title = paper.title
    doc.core_properties.author = "Портал ИВИТШ КГУ"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- PDF ---------------------------------------------------------------------------------------

def render_pdf(paper: Paper) -> bytes:
    size, leading = 14, 16.5
    base = ParagraphStyle("base", fontName="DocSerif", fontSize=size, leading=leading)
    address = ParagraphStyle("address", parent=base, leftIndent=_address_indent(paper.header) * cm)
    title = ParagraphStyle("title", parent=base, fontName="DocSerif-Bold", alignment=TA_CENTER, spaceBefore=36, spaceAfter=18)
    body = ParagraphStyle("body", parent=base, alignment=TA_JUSTIFY, firstLineIndent=1.25 * cm, leading=size * 1.5)
    item = ParagraphStyle("item", parent=body, alignment=TA_LEFT, leftIndent=_LIST_INDENT * cm, firstLineIndent=-_DASH_HANG * cm)
    right = ParagraphStyle("right", parent=base, alignment=TA_RIGHT)

    flow = [Paragraph(escape(line), address) for line in paper.header]
    flow.append(Paragraph(escape(paper.title), title))
    flow += [Paragraph(escape(text), body) for text in paper.paragraphs]
    flow += [Paragraph(_DASH + escape(text), item) for text in paper.items]
    flow += [Paragraph(escape(text), body) for text in paper.closing]
    flow.append(Spacer(1, 42))
    sign = Table(
        [[Paragraph(escape(_dated(paper.dated)), base), Paragraph(escape(f"____________ / {paper.signed}"), right)]],
        colWidths=[_TEXT_WIDTH / 2 * cm] * 2,
    )
    sign.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    flow.append(sign)

    out = io.BytesIO()
    SimpleDocTemplate(
        out, pagesize=A4, title=paper.title, author="Портал ИВИТШ КГУ",
        leftMargin=_MARGIN_LEFT * cm, rightMargin=_MARGIN_RIGHT * cm,
        topMargin=_MARGIN_Y * cm, bottomMargin=_MARGIN_Y * cm,
    ).build(flow)
    return out.getvalue()
